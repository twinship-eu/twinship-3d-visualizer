"""Builds the futuristic vessel's time-series template: one sheet, with sample data.

TIMESERIES: one row per moment, one column per signal, grouped by component.
  row 1  the component: its ID in the KG and its name
  row 2  the column's machine name, unique: <COMPONENT-ID>.TWINSHIP-<PARAMETER>-IN-<UNIT>
  row 3  the type
  row 4  the warning limit (a number: from it up; "<n": from it down); empty for none
  row 5  the alarm limit, the same way
  row 6+ the data, TIMESTAMP-UTC first

Usage: python generate-timeseries-template.py <output.xlsx>   (needs openpyxl)
"""

import math
import random
import sys
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

random.seed(7)
OUT = sys.argv[1]
START = datetime(2026, 9, 29, 6, 0, 0, tzinfo=timezone.utc)
STEP_S = 5
STEPS = 30 * 60 // STEP_S + 1  # 30 minutes, every 5 s
HEADER_ROWS = 5

GROUP_FILLS = ["1F3A5F", "2E5A88"]
NAME_FILL = PatternFill("solid", fgColor="DCE6F1")
META_FILL = PatternFill("solid", fgColor="F2F2F2")
LIMIT_FILL = PatternFill("solid", fgColor="FFF2CC")


class Walk:
    """A bounded random walk round a baseline: what a sensor's reading does."""

    def __init__(self, value, spread, pull=0.08):
        self.base, self.value, self.spread, self.pull = value, value, spread, pull

    def __call__(self):
        self.value += random.uniform(-1, 1) * self.spread * 0.5 + (self.base - self.value) * self.pull
        return self.value


# --- The simulated voyage, step by step: every column's value comes from here ---
state = {"lat": 56.12, "lon": 7.45, "dtg": 184.0, "fuel": {"FV-GN-1": 812.4, "FV-GN-2": 798.1}, "soc": 81.0}
walk = {
    "heading": Walk(87, 1.0), "sog": Walk(18.0, 0.2), "stw": Walk(18.6, 0.15), "wind": Walk(17.6, 0.8),
    "wind_dir": Walk(40, 3), "hs": Walk(2.0, 0.08), "draft_f": Walk(7.37, 0.01), "draft_a": Walk(7.49, 0.01),
    "cloud": Walk(35, 2), "ess": Walk(40, 30),
    **{f"{g}.{k}": w for g in ("FV-GN-1", "FV-GN-2") for k, w in
       (("power", Walk(5150, 80)), ("cool", Walk(78, 0.5)), ("avg", Walk(372, 3)))},
    **{f"{n}.{k}": w for n in (1, 2) for k, w in
       (("em", Walk(4800, 60)), ("rpm", Walk(140, 1.5)), ("pitch", Walk(96, 0.4)), ("rudder", Walk(1.1, 0.5)),
        ("wasp", Walk(220, 12)), ("saving", Walk(5, 0.2)), ("fan", Walk(1480, 15)))},
    **{f"SG-0{i}": Walk(v, 1.2) for i, v in ((1, 54), (2, 58), (3, 72), (4, 66), (5, 49))},
}
over_limit = []


def step(i):
    """Every column's value at step i."""
    t = START + timedelta(seconds=i * STEP_S)
    seconds, minute = i * STEP_S, i * STEP_S / 60
    heading, sog = walk["heading"](), walk["sog"]()
    state["lat"] += sog * STEP_S / 3600 / 60 * math.cos(math.radians(heading))
    state["lon"] += sog * STEP_S / 3600 / 60 * math.sin(math.radians(heading)) / math.cos(math.radians(state["lat"]))
    state["dtg"] -= sog * STEP_S / 3600
    v = {
        "TIMESTAMP-UTC": t.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "FV-RR-1.TWINSHIP-LATITUDE-IN-DEG": round(state["lat"], 5),
        "FV-RR-1.TWINSHIP-LONGITUDE-IN-DEG": round(state["lon"], 5),
        "FV-RR-1.TWINSHIP-HEADING-IN-DEG": round(heading % 360, 1),
        "FV-RR-1.TWINSHIP-COURSE-OVER-GROUND-IN-DEG": round((heading + 2) % 360, 1),
        "FV-RR-1.TWINSHIP-SPEED-OVER-GROUND-IN-KN": round(sog, 2),
        "FV-RR-1.TWINSHIP-SPEED-THROUGH-WATER-IN-KN": round(walk["stw"](), 2),
        "FV-RR-1.TWINSHIP-RATE-OF-TURN-IN-DEG_PER_MIN": round(random.uniform(-0.3, 0.5), 2),
        "FV-RR-1.TWINSHIP-DRAFT-FORE-IN-M": round(walk["draft_f"](), 2),
        "FV-RR-1.TWINSHIP-DRAFT-AFT-IN-M": round(walk["draft_a"](), 2),
        "FV-RR-1.TWINSHIP-ROLL-IN-DEG": round(1.6 * math.sin(2 * math.pi * seconds / 11.8) + random.uniform(-0.3, 0.3), 2),
        "FV-RR-1.TWINSHIP-PITCH-IN-DEG": round(0.5 * math.sin(2 * math.pi * seconds / 7.4 + 1) + random.uniform(-0.1, 0.1), 2),
        "FV-RR-1.TWINSHIP-HEAVE-IN-M": round(0.8 * math.sin(2 * math.pi * seconds / 7.4) + random.uniform(-0.1, 0.1), 2),
        "FV-RR-1.TWINSHIP-OPERATIONAL-MODE": "SAILING",
        "FV-RR-1.TWINSHIP-AUTONOMY-LEVEL": "L3",
        "FV-RR-1.TWINSHIP-DISTANCE-TO-GO-IN-NM": round(state["dtg"], 2),
        "FV-RR-1.TWINSHIP-TRUE-WIND-SPEED-IN-KN": round(walk["wind"](), 1),
        "FV-RR-1.TWINSHIP-TRUE-WIND-DIRECTION-IN-DEG": round(walk["wind_dir"]() % 360, 1),
        "FV-RR-1.TWINSHIP-SIGNIFICANT-WAVE-HEIGHT-IN-M": round(walk["hs"](), 2),
        "FV-RR-1.TWINSHIP-CLOUD-COVER-PERCENTAGE": round(min(max(walk["cloud"](), 0), 100)),
        "FV-RR-1.TWINSHIP-PRECIPITATION-IN-MM_PER_HR": 0,
    }
    for g in ("FV-GN-1", "FV-GN-2"):
        power = walk[f"{g}.power"]()
        fuel = power * 0.188  # ~188 g/kWh
        state["fuel"][g] += fuel * STEP_S / 3600 / 1000
        avg = walk[f"{g}.avg"]()
        # Genset 1's cylinder 4 heats from minute 12, over 400 C for a quarter of an hour, then eases back
        if g == "FV-GN-1" and 12 <= minute <= 32:
            hottest, cylinder = avg + 12 + 34 * math.sin(math.pi * min((minute - 12) / 20, 1)) + random.uniform(-2, 2), 4
            if hottest > 400:
                over_limit.append(t)
        else:
            hottest, cylinder = avg + 14 + random.uniform(-3, 3), random.choice([2, 7, 9])
        v.update({
            f"{g}.TWINSHIP-SHAFT-SPEED-IN-REV_PER_MIN": round(720 + random.uniform(-2, 2), 1),
            f"{g}.TWINSHIP-GENERATOR-POWER-IN-KW": round(power),
            f"{g}.TWINSHIP-LOAD-PERCENTAGE": round(100 * power / 6300, 1),
            f"{g}.TWINSHIP-FUEL-CONSUMPTION-RATE-IN-KG_PER_HR": round(fuel, 1),
            f"{g}.TWINSHIP-TOTAL-FUEL-CONSUMPTION-IN-MT": round(state["fuel"][g], 3),
            f"{g}.TWINSHIP-COOLING-WATER-TEMPERATURE-IN-°C": round(walk[f"{g}.cool"](), 1),
            f"{g}.TWINSHIP-EXHAUST-GAS-TEMPERATURE-MAX-IN-°C": round(hottest, 1),
            f"{g}.TWINSHIP-EXHAUST-GAS-TEMPERATURE-MAX-CYLINDER": cylinder,
        })
    ess = walk["ess"]()
    state["soc"] -= ess * STEP_S / 3600 / 7000 * 100
    v.update({
        "FV-ESS-1.TWINSHIP-ESS-STATE-OF-CHARGE-PERCENTAGE": round(state["soc"], 2),
        "FV-ESS-1.TWINSHIP-ESS-POWER-IN-KW": round(ess),
    })
    for n in (1, 2):
        power, rpm = walk[f"{n}.em"](), walk[f"{n}.rpm"]()
        wasp = walk[f"{n}.wasp"]()
        v.update({
            f"FV-EMSH-{n}.TWINSHIP-EMSH-POWER-IN-KW": round(power),
            f"FV-EMSH-{n}.TWINSHIP-EMSH-SPEED-REV-PER-MIN": round(rpm, 1),
            f"FV-EMSH-{n}.TWINSHIP-EMSH-TORQUE-kNm": round(power / (2 * math.pi * rpm / 60), 1),
            f"FV-CPP-{n}.TWINSHIP-CPP-PITCH-PERCENTAGE": round(walk[f"{n}.pitch"](), 1),
            f"FV-CPP-{n}.TWINSHIP-RUDDER-ANGLE-DEGREE": round(walk[f"{n}.rudder"](), 1),
            f"FV-WASP-{n}.TWINSHIP-WASP-POWER-kW": round(wasp),
            f"FV-WASP-{n}.TWINSHIP-WASP-POWER-REDUCTION-EFFICIENCY-PERCENTAGE": round(walk[f"{n}.saving"](), 1),
            f"FV-WASP-{n}.TWINSHIP-WASP-THRUST-IN-kN": round(wasp / 5.9, 1),
            f"FV-WASP-{n}.TWINSHIP-WASP-FAN-SPEED-IN-REV_PER_MIN": round(walk[f"{n}.fan"]()),
        })
    for i_gauge in range(1, 6):
        v[f"SG-0{i_gauge}.TWINSHIP-BENDING-STRESS-PERCENTAGE-OF-ALLOWABLE"] = round(walk[f"SG-0{i_gauge}"](), 1)
    return v


# --- The columns: component group, type, and limits ---------------------------
WEATHER_WORDS = ("WIND", "WAVE", "CLOUD", "PRECIPITATION")
GROUPS = [  # (label, which columns): in the sheet's order
    ("TIME", lambda column: column == "TIMESTAMP-UTC"),
    ("FV-RR-1 · VESSEL: NAVIGATION, MOTION, VOYAGE",
     lambda column: column.startswith("FV-RR-1.") and not any(word in column for word in WEATHER_WORDS)),
    ("FV-RR-1 · VESSEL: WEATHER", lambda column: column.startswith("FV-RR-1.")),
    ("FV-GN-1 · GENSET 1", lambda column: column.startswith("FV-GN-1.")),
    ("FV-GN-2 · GENSET 2", lambda column: column.startswith("FV-GN-2.")),
    ("FV-ESS-1 · ENERGY STORAGE", lambda column: column.startswith("FV-ESS-1.")),
    ("FV-EMSH-1 · MOTOR SHAFT 1", lambda column: column.startswith("FV-EMSH-1.")),
    ("FV-EMSH-2 · MOTOR SHAFT 2", lambda column: column.startswith("FV-EMSH-2.")),
    ("FV-CPP-1 · PROPELLER 1", lambda column: column.startswith("FV-CPP-1.")),
    ("FV-CPP-2 · PROPELLER 2", lambda column: column.startswith("FV-CPP-2.")),
    ("FV-WASP-1 · SUCTION SAIL 1", lambda column: column.startswith("FV-WASP-1.")),
    ("FV-WASP-2 · SUCTION SAIL 2", lambda column: column.startswith("FV-WASP-2.")),
    ("HULL STRAIN GAUGES (proposed)", lambda column: column.startswith("SG-")),
]
TYPES = {
    "TIMESTAMP-UTC": "DATETIME (ISO 8601, UTC)",
    "TWINSHIP-OPERATIONAL-MODE": "ENUM (SAILING | MANOEUVRING | IN-PORT)",
    "TWINSHIP-AUTONOMY-LEVEL": "ENUM (L0 … L4)",
    "TWINSHIP-EXHAUST-GAS-TEMPERATURE-MAX-CYLINDER": "INTEGER",
}
LIMITS = {  # (warning, alarm): a number from it up, "<n" from it down
    "TWINSHIP-ROLL-IN-DEG": (15, 25),
    "TWINSHIP-TRUE-WIND-SPEED-IN-KN": (35, 45),
    "TWINSHIP-SIGNIFICANT-WAVE-HEIGHT-IN-M": (4, 6),
    "TWINSHIP-LOAD-PERCENTAGE": (90, 100),
    "TWINSHIP-COOLING-WATER-TEMPERATURE-IN-°C": (88, 95),
    "TWINSHIP-EXHAUST-GAS-TEMPERATURE-MAX-IN-°C": (400, 420),
    "TWINSHIP-ESS-STATE-OF-CHARGE-PERCENTAGE": ("<25", "<15"),
    "TWINSHIP-EMSH-SPEED-REV-PER-MIN": (155, 165),
    "TWINSHIP-BENDING-STRESS-PERCENTAGE-OF-ALLOWABLE": (80, 90),
}

rows = [step(i) for i in range(STEPS)]
columns = list(rows[0].keys())


def group_of(column):
    """The index of the first group a column falls in."""
    return next(index for index, (_, matches) in enumerate(GROUPS) if matches(column))


order = {column: index for index, column in enumerate(columns)}
columns.sort(key=lambda column: (group_of(column), order[column]))

wb = Workbook()
readme = wb.active
readme.title = "README"
for line in [
    "TwinShip futuristic RoRo vessel — operational time series (template, sample data)",
    "",
    "Fill in the TIMESERIES sheet: one row per moment, one column per signal.",
    "Row 1: the component (its ID in the KG). Row 2: the column's name — keep it as it is, the viewer reads it.",
    "Row 3: the type. Rows 4-5: warning and alarm limits (a number: from it up; '<n': from it down). Data from row 6.",
    "TIMESTAMP-UTC: ISO 8601 in UTC, any regular interval (1-10 s reads smoothly). Leave a cell empty for no reading.",
    "Columns you have no data for can be deleted or left empty: the viewer simulates what the file does not give.",
    "Sample data: 30 min every 5 s under way; genset 1's cylinder 4 exhaust rises over 400 C for a quarter of an hour.",
    "Into the viewer: python scripts/convert-timeseries.py <this.xlsx> (see docs/live-data/futuristic-vessel-timeseries.md).",
]:
    readme.append([line])
readme.column_dimensions["A"].width = 120
readme["A1"].font = Font(bold=True, size=14)

ws = wb.create_sheet("TIMESERIES")
wb.active = 1
for index, column in enumerate(columns, start=1):
    parameter = column.split(".", 1)[-1]
    warning, alarm = LIMITS.get(parameter, (None, None))
    ws.cell(2, index, column)
    ws.cell(3, index, TYPES.get(parameter, "DOUBLE"))
    ws.cell(4, index, warning if index > 1 else "WARNING AT")
    ws.cell(5, index, alarm if index > 1 else "ALARM AT")
    ws.column_dimensions[get_column_letter(index)].width = 22 if index > 1 else 24
    for row_index, fill in ((2, NAME_FILL), (3, META_FILL), (4, LIMIT_FILL), (5, LIMIT_FILL)):
        cell = ws.cell(row_index, index)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(2, index).font = Font(bold=True, size=9)
    ws.cell(3, index).font = Font(italic=True, color="666666", size=9)
# Row 1: the groups, merged over their columns
start = 1
for index in range(1, len(columns) + 2):
    same = index <= len(columns) and group_of(columns[index - 1]) == group_of(columns[start - 1])
    if not same:
        group = group_of(columns[start - 1])
        ws.cell(1, start, GROUPS[group][0])
        if index - 1 > start:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=index - 1)
        cell = ws.cell(1, start)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GROUP_FILLS[group % 2])
        start = index
ws.row_dimensions[2].height = 60
for row in rows:
    ws.append([row[column] for column in columns])
ws.freeze_panes = ws.cell(HEADER_ROWS + 1, 2)

wb.save(OUT)
print(f"wrote {OUT}: {len(columns)} columns, {len(rows)} rows; cylinder 4 over 400 C from {over_limit[0]:%H:%M:%S} to {over_limit[-1]:%H:%M:%S}")
