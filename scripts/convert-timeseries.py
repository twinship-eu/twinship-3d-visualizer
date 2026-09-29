"""Converts the vessel's time-series workbook into the recording the 3D viewer plays.

The viewer never reads Excel: this runs beforehand, the way a data pipeline
would, and writes JSON — what a live stream would deliver.

Input: the TIMESERIES sheet of a workbook laid out as
docs/live-data/futuristic-vessel-timeseries-template.xlsx:
  row 2  column names, <COMPONENT-ID>.TWINSHIP-<PARAMETER>-IN-<UNIT>
  row 4  warning limits (n: from it up; "<n": from it down)
  row 5  alarm limits
  row 6+ data, TIMESTAMP-UTC first

Output (column-oriented, compact): {
  "source", "convertedAt",
  "columns": [{"name", "component", "parameter", "warning", "alarm"}],  limits as {"value", "direction"}
  "times": [ms since the epoch, ...],
  "values": {name: [number | string | null, ...]}
}

Usage:
  python scripts/convert-timeseries.py <workbook.xlsx> [output.json]
Output defaults to public/live-data/futuristic-vessel-timeseries.json, the
recording the viewer plays. For data that must not be committed, write
public/live-data/<name>.local.json (git ignores those) and point the viewer at
it in .env.local: NEXT_PUBLIC_LIVE_RECORDING_URL=/live-data/<name>.local.json
"""

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SHEET = "TIMESERIES"
NAMES_ROW, WARNING_ROW, ALARM_ROW, FIRST_DATA_ROW = 2, 4, 5, 6
TIME_COLUMN = "TIMESTAMP-UTC"
DEFAULT_OUTPUT = Path(__file__).resolve().parent.parent / "public" / "live-data" / "futuristic-vessel-timeseries.json"


def limit(cell, where):
    """A limit cell as {value, direction}; None when empty."""
    if cell is None or cell == "":
        return None
    text = str(cell).strip()
    try:
        if text.startswith("<"):
            return {"value": float(text[1:]), "direction": "low"}
        return {"value": float(text), "direction": "high"}
    except ValueError:
        raise SystemExit(f"{SHEET}!{where}: a limit is a number, or <number for a low one; found {cell!r}")


def moment(cell, where):
    """A timestamp cell, ISO 8601 text or an Excel date, as ms since the epoch (UTC)."""
    if isinstance(cell, datetime):
        stamp = cell if cell.tzinfo else cell.replace(tzinfo=timezone.utc)
    else:
        try:
            stamp = datetime.fromisoformat(str(cell).strip().replace("Z", "+00:00"))
        except ValueError:
            raise SystemExit(f"{SHEET}!{where}: not a UTC timestamp like 2026-09-29T06:00:00Z; found {cell!r}")
        if stamp.tzinfo is None:
            stamp = stamp.replace(tzinfo=timezone.utc)
    return round(stamp.timestamp() * 1000)


def value(cell):
    """A data cell: a number, text for the enums, or None for no reading."""
    if cell is None or cell == "":
        return None
    if isinstance(cell, (int, float)):
        return cell
    text = str(cell).strip()
    try:
        return float(text)
    except ValueError:
        return text


def convert(workbook_path, output_path):
    sheet = load_workbook(workbook_path, read_only=True, data_only=True)[SHEET]
    rows = list(sheet.iter_rows(values_only=True))
    names = [str(name).strip() if name else None for name in rows[NAMES_ROW - 1]]
    if names[0] != TIME_COLUMN:
        raise SystemExit(f"{SHEET}: column A of row {NAMES_ROW} must be {TIME_COLUMN}, found {names[0]!r}")

    columns = []
    for index, name in enumerate(names[1:], start=1):
        if not name:
            continue
        component, _, parameter = name.partition(".")
        columns.append({
            "name": name,
            "component": component,
            "parameter": parameter,
            "warning": limit(rows[WARNING_ROW - 1][index], f"{get_column_letter(index + 1)}{WARNING_ROW}"),
            "alarm": limit(rows[ALARM_ROW - 1][index], f"{get_column_letter(index + 1)}{ALARM_ROW}"),
            "index": index,
        })

    times, values = [], {column["name"]: [] for column in columns}
    for number, row in enumerate(rows[FIRST_DATA_ROW - 1:], start=FIRST_DATA_ROW):
        if not row or row[0] in (None, ""):
            continue
        stamp = moment(row[0], f"A{number}")
        if times and stamp <= times[-1]:
            raise SystemExit(f"{SHEET}!A{number}: {TIME_COLUMN} must increase from row to row")
        times.append(stamp)
        for column in columns:
            values[column["name"]].append(value(row[column["index"]] if column["index"] < len(row) else None))
    if len(times) < 2:
        raise SystemExit(f"{SHEET}: at least two rows of data are needed, found {len(times)}")

    for column in columns:
        del column["index"]
    recording = {
        "source": Path(workbook_path).name,
        "convertedAt": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "columns": columns,
        "times": times,
        "values": values,
    }
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    Path(output_path).write_text(json.dumps(recording, separators=(",", ":"), ensure_ascii=False))
    span = (times[-1] - times[0]) / 60000
    print(f"{output_path}: {len(columns)} columns, {len(times)} rows, {span:.1f} min")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    convert(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUTPUT)
